"""
store.py  —  JSON tracker + Excel sync for Telegram-submitted reels.
================================================================================
The JSON store (telegram_bot/telegram_reels.json) is the dedup / timeline layer
for reels the owner drops into the Telegram chat:

    {
      "last_update_id": 0,
      "reels": {
        "<shortcode>": {
          "url": "...", "received_date": "...", "status": "RECEIVED",
          "telegram_message_id": 42, "message_text": "...",
          "draft_requested": false, "draft_count": 0,
          "processed_date": null, "analyzed_date": null
        }
      }
    }

Brain/Reels_Log.xlsx stays the spreadsheet of truth for the wider pipeline —
sync_to_excel() appends a Category=NICHE / Status=SCRAPED / Source=TELEGRAM row
so the normal process → transcribe → analyze chain can pick the reel up.

Usage:
    uv run telegram_bot/store.py            # dump store stats
    uv run telegram_bot/store.py --pending  # list RECEIVED reels as JSON
"""

import os
import re
import sys
import json
import time
import argparse
import tempfile
import contextlib
import datetime
from urllib.parse import urlparse

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import pandas as pd
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "System")))
from workspace_lock import atomic_path, locked

try:                                     # package import (from telegram_bot import store)
    from . import config
except ImportError:                      # script import (uv run telegram_bot/store.py)
    import config                        # type: ignore

STORE_PATH       = config.STORE_PATH
SPREADSHEET_PATH = config.SPREADSHEET_PATH

# Small dedicated advisory lock for the JSON store's load->mutate->save cycle,
# so a concurrent fetch / status update can never silently drop records.
STORE_LOCK_PATH  = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                ".telegram_store.lock")

# Column order used by System/scrape.py + System/process_reels.py.
# "Source" is appended LAST on purpose: those scripts address Category/Status by
# positional index (row[1] / row[13]), so extra columns must never shift them.
BASE_COLUMNS = [
    "Shortcode", "Category", "Creator Name", "Username", "Reel Link", "Video URL",
    "Publication Date", "Like Count", "Comments Count", "Shares Count",
    "Caption", "Topic/Summary", "Scraped Date", "Status",
    "Processed Date", "Brain Note Path",
]
SOURCE_COLUMN = "Source"
COLUMNS = BASE_COLUMNS + [SOURCE_COLUMN]

# Statuses used inside the JSON store
STATUS_RECEIVED = "RECEIVED"

# Draft count used when the message asks for drafts without naming a number
DEFAULT_DRAFT_COUNT = 2
MAX_DRAFT_COUNT = 10

REEL_URL_RE = re.compile(
    r"https?://(?:www\.)?instagram\.com/(?:reel|reels|p)/[A-Za-z0-9_-]+",
    re.IGNORECASE,
)


def _now() -> str:
    return datetime.datetime.now().isoformat(timespec="seconds")


@contextlib.contextmanager
def store_lock(timeout: float = 60.0):
    """Serialize JSON-store read-modify-write transactions across processes.

    Small advisory file lock (stdlib only — no broad dependency). Fetch cycles
    and status updates run their whole load -> mutate -> save cycle under it,
    so two simultaneous writers can never overwrite each other's records.
    """
    os.makedirs(os.path.dirname(STORE_LOCK_PATH), exist_ok=True)
    with open(STORE_LOCK_PATH, "a+b") as handle:
        handle.seek(0, os.SEEK_END)
        if handle.tell() == 0:
            handle.write(b"0")
            handle.flush()

        deadline = time.monotonic() + timeout
        acquired = False
        while time.monotonic() < deadline:
            try:
                handle.seek(0)
                if os.name == "nt":
                    import msvcrt
                    msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                acquired = True
                break
            except (OSError, BlockingIOError):
                time.sleep(0.1)

        if not acquired:
            raise TimeoutError(
                f"timed out acquiring Telegram store lock: {STORE_LOCK_PATH}")

        try:
            yield
        finally:
            handle.seek(0)
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


# ── JSON store ───────────────────────────────────────────────────────────────
def _empty_store() -> dict:
    return {"last_update_id": 0, "reels": {}}


def load_store() -> dict:
    """Load the JSON store. Returns a fresh structure if the file is absent
    or unreadable (nothing is written to disk here — call save_store())."""
    if not os.path.exists(STORE_PATH):
        return _empty_store()
    try:
        with open(STORE_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError) as e:
        print(f"[WARN] Could not read {STORE_PATH}: {e} — starting a fresh store",
              file=sys.stderr)
        return _empty_store()
    if not isinstance(data, dict):
        return _empty_store()
    data.setdefault("last_update_id", 0)
    data.setdefault("reels", {})
    return data


def save_store(data: dict) -> None:
    """Atomically write the JSON store."""
    os.makedirs(os.path.dirname(STORE_PATH), exist_ok=True)
    tmp = f"{STORE_PATH}.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    os.replace(tmp, STORE_PATH)


# ── URL / shortcode helpers ──────────────────────────────────────────────────
def extract_shortcode(url: str) -> str:
    """Extract the shortcode from an Instagram URL.

    Mirrors System/ingest_reel.py:get_video_id but also handles /reels/ and
    returns "" (not "unknown_video") when nothing can be parsed."""
    if not url:
        return ""
    parsed = urlparse(url.strip())
    parts = [p for p in parsed.path.strip("/").split("/") if p]
    for key in ("reel", "reels", "p"):
        if key in parts:
            idx = parts.index(key)
            if idx + 1 < len(parts):
                return parts[idx + 1]
    m = re.search(r"/(?:reel|reels|p)/([A-Za-z0-9_-]+)", url)
    return m.group(1) if m else ""


def find_reel_urls(text: str) -> list[str]:
    """All Instagram reel/post URLs in `text`, de-duplicated, order preserved."""
    if not text:
        return []
    return list(dict.fromkeys(REEL_URL_RE.findall(text)))


def canonical_url(shortcode: str) -> str:
    return f"https://www.instagram.com/reel/{shortcode}/"


# ── Excel lookups ────────────────────────────────────────────────────────────
_excel_cache: set[str] | None = None


def excel_shortcodes(refresh: bool = False) -> set[str]:
    """Shortcodes already present in Brain/Reels_Log.xlsx (cached per process)."""
    global _excel_cache
    if _excel_cache is not None and not refresh:
        return _excel_cache
    codes: set[str] = set()
    if os.path.exists(SPREADSHEET_PATH):
        try:
            df = pd.read_excel(SPREADSHEET_PATH)
            if "Shortcode" in df.columns:
                codes = set(df["Shortcode"].dropna().astype(str).str.strip())
        except Exception as e:
            print(f"[WARN] Could not read {SPREADSHEET_PATH}: {e}", file=sys.stderr)
    _excel_cache = codes
    return codes


def is_duplicate(shortcode: str, store: dict | None = None) -> bool:
    """True when the shortcode is already tracked in the JSON store OR the Excel log."""
    if not shortcode:
        return False
    data = store if store is not None else load_store()
    if shortcode in data.get("reels", {}):
        return True
    return shortcode in excel_shortcodes()


def duplicate_reason(shortcode: str, store: dict | None = None) -> str:
    """Human-readable explanation for a skipped reel ("" when not a duplicate)."""
    data = store if store is not None else load_store()
    entry = data.get("reels", {}).get(shortcode)
    if entry:
        return f"already in store as {entry.get('status', STATUS_RECEIVED)}"
    if shortcode in excel_shortcodes():
        return "already in Brain/Reels_Log.xlsx"
    return ""


# ── Draft-intent parsing ─────────────────────────────────────────────────────
_NUMBER_WORDS = {
    "a": 1, "an": 1, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5,
    "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
}
_NUM = r"\d{1,2}|" + "|".join(_NUMBER_WORDS)
_TARGET = r"(?:re)?drafts?|scripts?|versions?|variations?"
_VERB = r"(?:re)?draft|create|make|write|generate|give|need|want|do"

# "2 scripts", "three quick drafts"  → number BEFORE the noun
_COUNT_BEFORE_RE = re.compile(
    rf"\b({_NUM})\b(?:\s+\w+){{0,2}}?\s+\b(?:{_TARGET})\b", re.IGNORECASE)
# "draft 2", "give me three"         → number AFTER the verb
_COUNT_AFTER_RE = re.compile(
    rf"\b(?:{_VERB})\b(?:\s+(?:me|us))?\s+({_NUM})\b", re.IGNORECASE)
# Does the message ask for scripts at all?
_INTENT_RE = re.compile(rf"\b(?:{_TARGET})\b", re.IGNORECASE)


def _to_int(token: str) -> int:
    token = token.strip().lower()
    if token.isdigit():
        return int(token)
    return _NUMBER_WORDS.get(token, DEFAULT_DRAFT_COUNT)


def detect_draft_intent(message_text: str) -> tuple[bool, int]:
    """
    Parse a message for a script/draft request.

    Handles (case-insensitive, flexible word order):
        "draft 2 scripts based on this reel https://…"  -> (True, 2)
        "draft 3 scripts https://…"                     -> (True, 3)
        "create 2 drafts from this https://…"           -> (True, 2)
        "make me three scripts"                         -> (True, 3)
        "draft scripts from this"                       -> (True, DEFAULT_DRAFT_COUNT)
        "https://… " (bare link)                        -> (False, 0)
    """
    if not message_text:
        return (False, 0)

    # Strip URLs first so digits inside links never read as a count.
    text = REEL_URL_RE.sub(" ", message_text)
    text = re.sub(r"https?://\S+", " ", text)

    if not _INTENT_RE.search(text):
        return (False, 0)

    for pattern in (_COUNT_BEFORE_RE, _COUNT_AFTER_RE):
        m = pattern.search(text)
        if m:
            count = _to_int(m.group(1))
            if count >= 1:
                return (True, min(count, MAX_DRAFT_COUNT))

    return (True, DEFAULT_DRAFT_COUNT)


# ── Store mutations ──────────────────────────────────────────────────────────
def add_reel(url: str, message_id: int, message_text: str,
             store: dict | None = None, save: bool | None = None) -> dict:
    """
    Add a reel to the JSON store and return its entry.

    Pass `store` to batch several adds in memory (then save_store() once);
    with no `store` the file is loaded and written for you.
    """
    owns_store = store is None
    data = store if store is not None else load_store()
    if save is None:
        save = owns_store

    shortcode = extract_shortcode(url)
    if not shortcode:
        raise ValueError(f"Could not extract a shortcode from URL: {url!r}")

    draft_requested, draft_count = detect_draft_intent(message_text)

    entry = {
        "url": url,
        "received_date": _now(),
        "status": STATUS_RECEIVED,
        "telegram_message_id": message_id,
        "message_text": message_text,
        "draft_requested": draft_requested,
        "draft_count": draft_count,
        "processed_date": None,
        "analyzed_date": None,
    }
    data.setdefault("reels", {})[shortcode] = entry

    if save:
        save_store(data)
    return entry


def update_reel_status(shortcode: str, status: str,
                       store: dict | None = None, save: bool | None = None) -> dict | None:
    """Update a reel's status (also stamps processed/analyzed dates)."""
    owns_store = store is None
    data = store if store is not None else load_store()
    if save is None:
        save = owns_store

    entry = data.get("reels", {}).get(shortcode)
    if entry is None:
        print(f"[WARN] {shortcode} not found in the store — status not updated",
              file=sys.stderr)
        return None

    status = status.upper()
    entry["status"] = status
    if status == "PROCESSED":
        entry["processed_date"] = _now()
    elif status == "ANALYZED":
        entry["analyzed_date"] = _now()

    if save:
        save_store(data)
    return entry


def update_last_update_id(update_id: int,
                          store: dict | None = None, save: bool | None = None) -> int:
    """Persist the getUpdates offset (never moves backwards)."""
    owns_store = store is None
    data = store if store is not None else load_store()
    if save is None:
        save = owns_store

    current = int(data.get("last_update_id", 0) or 0)
    data["last_update_id"] = max(current, int(update_id))

    if save:
        save_store(data)
    return data["last_update_id"]


def get_pending_reels(store: dict | None = None) -> list:
    """Reels still at status RECEIVED, each entry annotated with its shortcode."""
    data = store if store is not None else load_store()
    pending = []
    for shortcode, entry in data.get("reels", {}).items():
        if str(entry.get("status", "")).upper() == STATUS_RECEIVED:
            item = dict(entry)
            item["shortcode"] = shortcode
            pending.append(item)
    return pending


# ── Excel sync ───────────────────────────────────────────────────────────────
def _load_sheet() -> pd.DataFrame:
    if os.path.exists(SPREADSHEET_PATH):
        try:
            df = pd.read_excel(SPREADSHEET_PATH)
        except Exception as e:
            raise RuntimeError(f"Could not read {SPREADSHEET_PATH}: {e}") from e
    else:
        df = pd.DataFrame(columns=BASE_COLUMNS)
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""
    # Keep known columns first (stable indices), then anything else we found.
    extra = [c for c in df.columns if c not in COLUMNS]
    return df[COLUMNS + extra]


def _save_sheet(df: pd.DataFrame) -> None:
    """Write the sheet back with the same styling System/*.py applies.

    Written to a sibling temp file with an .xlsx suffix (pandas/openpyxl
    reject other extensions) and atomically replaced, so a failed write
    never truncates a valid workbook.
    """
    from openpyxl.styles import Font, PatternFill, Alignment

    os.makedirs(os.path.dirname(SPREADSHEET_PATH), exist_ok=True)
    columns = list(df.columns)
    cat_idx = columns.index("Category") if "Category" in columns else None
    st_idx = columns.index("Status") if "Status" in columns else None

    directory = os.path.dirname(SPREADSHEET_PATH)
    fd, temp_path = tempfile.mkstemp(prefix=".reels_log_", suffix=".xlsx", dir=directory)
    os.close(fd)
    try:
        with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Reels Log")
            ws = writer.sheets["Reels Log"]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3864")
                cell.alignment = Alignment(horizontal="center")
            ws.freeze_panes = "A2"
            for row in ws.iter_rows(min_row=2):
                if cat_idx is not None and cat_idx < len(row):
                    cat_cell = row[cat_idx]
                    val = str(cat_cell.value).upper()
                    if val == "BRANDING":
                        cat_cell.fill = PatternFill("solid", fgColor="FFE699")
                    elif val == "NICHE":
                        cat_cell.fill = PatternFill("solid", fgColor="C6EFCE")
                if st_idx is not None and st_idx < len(row):
                    st_cell = row[st_idx]
                    val = str(st_cell.value).upper()
                    if val in ("PROCESSED", "ANALYZED"):
                        st_cell.fill = PatternFill("solid", fgColor="C6EFCE")
                    elif val == "SCRAPED":
                        st_cell.fill = PatternFill("solid", fgColor="FFEB9C")
                    elif val == "FAILED":
                        st_cell.fill = PatternFill("solid", fgColor="FFC7CE")
            for col in ws.columns:
                max_len = max((len(str(cell.value or "")) for cell in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 80)
        os.replace(temp_path, SPREADSHEET_PATH)
    except Exception:
        try:
            os.remove(temp_path)
        except OSError:
            pass
        raise


def _telegram_row(shortcode: str, url: str, message_text: str) -> dict:
    caption = (message_text or "").strip()
    summary = caption[:120] + "..." if len(caption) > 120 else caption
    return {
        "Shortcode":        shortcode,
        "Category":         "NICHE",
        "Creator Name":     "",
        "Username":         "",
        "Reel Link":        url or canonical_url(shortcode),
        "Video URL":        "",
        "Publication Date": "",
        "Like Count":       "",
        "Comments Count":   "",
        "Shares Count":     "",
        "Caption":          caption,
        "Topic/Summary":    summary,
        "Scraped Date":     datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "Status":           "SCRAPED",
        "Processed Date":   "",
        "Brain Note Path":  "",
        SOURCE_COLUMN:      "TELEGRAM",
    }


def _apply_telegram_rows(df: pd.DataFrame, reels: list[dict]) -> tuple[pd.DataFrame, dict]:
    """Mutate `df` for every Telegram reel (no I/O).

    New shortcodes get a Source=TELEGRAM row appended; known shortcodes only
    get Source=TELEGRAM stamped. Returns (df, {shortcode: result}).
    """
    results: dict[str, str] = {}
    new_rows: list[dict] = []
    for item in reels:
        shortcode = str(item.get("shortcode", "") or "").strip()
        url  = str(item.get("url", "") or "")
        text = str(item.get("message_text", "") or "")
        if not shortcode:
            results[shortcode] = "failed: missing shortcode"
            continue
        match = df["Shortcode"].astype(str).str.strip() == shortcode
        if match.any():
            already = (df.loc[match, SOURCE_COLUMN].astype(str).str.upper() == "TELEGRAM").all()
            df.loc[match, SOURCE_COLUMN] = "TELEGRAM"
            results[shortcode] = "unchanged" if already else "updated"
        else:
            new_rows.append(_telegram_row(shortcode, url, text))
            results[shortcode] = "added"
    if new_rows:
        appended = pd.DataFrame(new_rows)
        # Align with the full sheet column set so existing extras survive.
        appended = appended.reindex(columns=df.columns, fill_value="")
        df = pd.concat([df, appended], ignore_index=True)
    return df, results


@locked
def sync_to_excel(shortcode: str, url: str, message_text: str) -> str:
    """
    Add (or tag) the reel in Brain/Reels_Log.xlsx.

    New shortcode  -> append a row: Category=NICHE, Status=SCRAPED, Source=TELEGRAM.
    Known shortcode-> only stamp Source=TELEGRAM (never duplicates a row).

    Returns "added" | "updated" | "unchanged".
    """
    if not shortcode:
        raise ValueError("sync_to_excel() requires a shortcode")
    df = _load_sheet()
    df, results = _apply_telegram_rows(df, [{
        "shortcode": shortcode, "url": url, "message_text": message_text,
    }])
    _save_sheet(df)
    excel_shortcodes(refresh=True)
    return results.get(str(shortcode).strip(), "failed")


@locked
def sync_to_excel_batch(reels: list[dict]) -> dict:
    """
    Add (or tag) every Telegram reel in ONE workbook read/stylize/write.

    `reels` is a list of {"shortcode", "url", "message_text"} dicts. The sheet
    is loaded, updated and written once for the whole batch instead of once per
    reel. Returns {shortcode: "added" | "updated" | "unchanged"} for every reel
    so callers can still report per-shortcode `excel` status after one write.
    """
    if not reels:
        return {}
    df = _load_sheet()
    df, results = _apply_telegram_rows(df, reels)
    _save_sheet(df)
    excel_shortcodes(refresh=True)
    return results


# ── CLI (debug helper) ───────────────────────────────────────────────────────
def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect the Telegram reel store.")
    parser.add_argument("--pending", action="store_true",
                        help="Print RECEIVED reels as JSON")
    args = parser.parse_args()

    data = load_store()
    if args.pending:
        print(json.dumps(get_pending_reels(data), indent=2, ensure_ascii=False))
        return

    reels = data.get("reels", {})
    counts: dict[str, int] = {}
    for entry in reels.values():
        key = str(entry.get("status", "?")).upper()
        counts[key] = counts.get(key, 0) + 1
    print(json.dumps({
        "store_path": STORE_PATH,
        "exists": os.path.exists(STORE_PATH),
        "last_update_id": data.get("last_update_id", 0),
        "total_reels": len(reels),
        "by_status": counts,
        "excel_rows": len(excel_shortcodes()),
    }, indent=2))


if __name__ == "__main__":
    main()
